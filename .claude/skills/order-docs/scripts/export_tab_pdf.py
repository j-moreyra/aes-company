#!/usr/bin/env python3
"""Export one worksheet of a CS workbook to PDF, the way the quote skill does it.

    python3 export_tab_pdf.py --workbook CS_260.19034.xlsx --sheet "PO TO KUKLA" \
        --out "PO 330_07282026 KUKLA.pdf" [--print-area B1:H45] [--fit-height 1]

Copies the workbook, deletes every sheet except the target, sets the print area
and fit, then converts with LibreOffice. Deleting the other sheets matters: any
cross-sheet reference left behind renders as #NAME? once they are gone, which is
also why values written into a quote/PO/OA tab must be literals, never formulas
pointing at Profit Calc.

Exit codes: 0 written, 2 LibreOffice unavailable (caller should fall back).
"""
import argparse, os, shutil, subprocess, sys, tempfile


def soffice():
    for n in ("soffice", "libreoffice"):
        p = shutil.which(n)
        if p:
            return p
    for p in ("/Applications/LibreOffice.app/Contents/MacOS/soffice",
              "/usr/bin/soffice", "/usr/lib/libreoffice/program/soffice"):
        if os.path.exists(p):
            return p
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--sheet", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--print-area", default=None,
                    help="e.g. B1:H45. Omit to let LibreOffice decide.")
    ap.add_argument("--fit-height", type=int, default=1,
                    help="1 = force one page tall (default). 0 = fit width only, "
                         "let it paginate naturally — use for long documents.")
    a = ap.parse_args()

    exe = soffice()
    if not exe:
        print("LibreOffice not found — cannot export the workbook tab natively.",
              file=sys.stderr)
        print("Install with: brew install --cask libreoffice", file=sys.stderr)
        sys.exit(2)

    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as td:
        tmp_xlsx = os.path.join(td, "export.xlsx")
        shutil.copy2(a.workbook, tmp_xlsx)
        wb = load_workbook(tmp_xlsx)
        if a.sheet not in wb.sheetnames:
            print(f"sheet {a.sheet!r} not found; tabs: {wb.sheetnames}", file=sys.stderr)
            sys.exit(1)
        for s in [s for s in wb.sheetnames if s != a.sheet]:
            del wb[s]
        ws = wb[a.sheet]
        if a.print_area:
            ws.print_area = a.print_area
        ws.page_setup.orientation = "portrait"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = a.fit_height   # 0 = paginate naturally
        wb.save(tmp_xlsx)

        subprocess.run([exe, "--headless", "--calc", "--convert-to", "pdf",
                        "--outdir", td, tmp_xlsx],
                       check=True, capture_output=True, timeout=300)
        made = os.path.join(td, "export.pdf")
        if not os.path.exists(made):
            print("LibreOffice produced no PDF", file=sys.stderr)
            sys.exit(1)
        os.makedirs(os.path.dirname(os.path.abspath(a.out)) or ".", exist_ok=True)
        shutil.move(made, a.out)

    print(f"wrote {a.out} ({os.path.getsize(a.out):,} bytes, via LibreOffice)")


if __name__ == "__main__":
    main()
