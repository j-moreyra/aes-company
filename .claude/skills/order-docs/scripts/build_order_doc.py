#!/usr/bin/env python3
"""Render an AES Order Acknowledgement or Purchase Order from a JSON spec.

Writes the values into the matching tab of the project CS_*.xlsx workbook and
renders a PDF that mirrors the workbook's print layout.

    python3 build_order_doc.py spec.json [--no-xlsx] [--no-pdf] [--out DIR]

The JSON spec is produced by the model after reading the source documents; this
script does no interpretation. See SKILL.md for the schema.
"""
import argparse, base64, json, os, shutil, subprocess, sys, tempfile

CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, "..", "assets", "aes_logo.png")

AES_NAME = "ADVANCED ENGINEERING SYSTEMS, LLC"
AES_ADDR = "9380 SW 108 Street - Miami, Florida 33176"
AES_CONTACT = "Ph. (305) 596-1128   E-mail: info@advengsys.com"
AES_FEI = "FEI 26-3685134"


CURRENCY = {"USD": "$", "EUR": "€"}


def money(v):
    if v is None or v == "":
        return ""
    if isinstance(v, str):
        return v
    return f"{v:,.2f}"


def sym(spec):
    """Currency symbol. OA is billed to the client in USD; the PO to Kukla is
    placed in the currency of Kukla's quote, which is EUR."""
    cur = spec.get("currency") or ("USD" if spec.get("kind") == "oa" else "EUR")
    if cur not in CURRENCY:
        raise SystemExit(f"unknown currency {cur!r}; expected one of {list(CURRENCY)}")
    return CURRENCY[cur]


def esc(s):
    return ("" if s is None else str(s)).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def addr_block(lines):
    return "".join(f"<div>{esc(l)}</div>" for l in lines if l)


# ---------------------------------------------------------------- validation
def validate(spec):
    """Arithmetic and cross-document checks. Returns list of problems."""
    errs, warns = [], []
    kind = spec.get("kind")
    if kind not in ("oa", "po_kukla"):
        errs.append(f"kind must be 'oa' or 'po_kukla', got {kind!r}")
    items = spec.get("items") or []
    if not items:
        errs.append("no line items")
    sub = 0.0
    for i, it in enumerate(items, 1):
        if it.get("heading"):          # a Kukla main-line title carries no price
            continue
        q, p = it.get("qty"), it.get("unit_price")
        if q is None or p is None:
            errs.append(f"item {i}: missing qty or unit_price"); continue
        line = round(q * p, 2)
        if it.get("total") is not None and abs(line - it["total"]) > 0.01:
            errs.append(f"item {i} ({it.get('description','?')}): {q} x {p} = {line}, spec says {it['total']}")
        it["total"] = line
        sub += line
    sub = round(sub, 2)
    if spec.get("subtotal") is not None and abs(sub - spec["subtotal"]) > 0.01:
        errs.append(f"subtotal mismatch: lines sum to {sub}, spec says {spec['subtotal']}")
    spec["subtotal"] = sub
    freight = spec.get("freight_amount")
    total = sub + (freight if isinstance(freight, (int, float)) else 0)
    if spec.get("total") is not None and abs(total - spec["total"]) > 0.01:
        errs.append(f"total mismatch: {sub} + freight = {total}, spec says {spec['total']}")
    spec["total"] = round(total, 2)

    # cross-check against the source document the model transcribed
    src = spec.get("source_totals") or {}
    if "subtotal" in src and abs(src["subtotal"] - sub) > 0.01:
        errs.append(f"line items sum to {sub} but the source document subtotal is {src['subtotal']} "
                    f"— reconcile before sending")
    if kind == "oa" and spec.get("show_subtotal", True) and not isinstance(freight, (int, float)):
        warns.append("OA has no numeric freight; the client PO usually excludes freight so AES adds it")
    return errs, warns


# ---------------------------------------------------------------- xlsx
def write_tab(spec, workbook, tab):
    from openpyxl import load_workbook
    wb = load_workbook(workbook)
    if tab not in wb.sheetnames:
        raise SystemExit(f"tab {tab!r} not in {os.path.basename(workbook)}; tabs: {wb.sheetnames}")
    ws = wb[tab]
    m = spec["cellmap"]

    def put(ref, val):
        if ref and val is not None:
            ws[ref] = val

    for ref, val in (m.get("fields") or {}).items():
        put(ref, val)
    rows = m["items"]
    r = rows["first_row"]
    for idx, it in enumerate(spec["items"]):
        rr = r + idx
        put(f"{rows['line']}{rr}", idx + 1)
        put(f"{rows['qty']}{rr}", it["qty"])
        if rows.get("item_no"):
            put(f"{rows['item_no']}{rr}", it.get("item_no"))
        put(f"{rows['description']}{rr}", it["description"])
        put(f"{rows['unit_price']}{rr}", it["unit_price"])
        # The template carries =qty*price in the total column and =SUM()/=+ in the
        # totals block. Only write a literal total when the caller asks for it,
        # otherwise leave the workbook's own formulas to compute.
        if rows.get("write_line_totals"):
            put(f"{rows['total']}{rr}", it["total"])
    for ref, val in (m.get("totals") or {}).items():
        put(ref, val)
    wb.save(workbook)
    return workbook


# ---------------------------------------------------------------- pdf
def html_for(spec):
    logo = ""
    if os.path.exists(LOGO):
        logo = ("data:image/png;base64,"
                + base64.b64encode(open(LOGO, "rb").read()).decode())
    oa = spec["kind"] == "oa"
    cur = sym(spec)
    title = "Order Acknowledgement" if oa else "Purchase Order"
    meta = spec["meta"]
    rows = []
    for i, it in enumerate(spec["items"], 1):
        # Kukla quotes use their own line numbering (110, 115, 130 …). Carry it
        # through so their sales desk can match the PO to the offer; fall back to
        # a simple 1..N sequence for client-facing documents.
        line_no = it.get("line_no", i)
        item_cell = f"<td class=c>{esc(it.get('item_no',''))}</td>" if not oa else ""
        if it.get("heading"):
            span = 6 if not oa else 5
            rows.append(f"<tr><td class=c>{esc(line_no)}</td>"
                        f"<td colspan={span-1}><b>{esc(it['description'])}</b></td></tr>")
            continue
        rows.append(
            f"<tr><td class=c>{esc(line_no)}</td><td class=c>{esc(it['qty'])}</td>{item_cell}"
            f"<td>{esc(it['description'])}</td>"
            f"<td class=n>{cur}&nbsp;&nbsp;{money(it['unit_price'])}</td>"
            f"<td class=n>{cur}&nbsp;&nbsp;{money(it['total'])}</td></tr>")
    blank = max(0, 18 - len(spec["items"]))
    ncols = 6 if not oa else 5
    rows += [f"<tr class=blank>{'<td></td>'*ncols}</tr>"] * blank
    metarows = "".join(
        f"<tr><td class=ml>{esc(k)}</td><td class=mv>{esc(v)}</td></tr>" for k, v in meta)
    ship = spec.get("ship_to")
    freight_label = spec.get("freight_label") or "FREIGHT"
    freight_val = spec.get("freight_amount")
    freight_txt = money(freight_val) if isinstance(freight_val, (int, float)) else esc(freight_val or "TBD")
    total_label = spec.get("total_label") or ("TOTAL DUE US" if oa else "TOTAL PURCHASE ORDER")
    hdr_item = "" if oa else "<th>ITEM</th>"
    hdr_total = "TOTAL US$" if oa else "TOTAL"

    # Some documents show subtotal + freight + total; others carry freight as a
    # numbered line item and show a single total. show_subtotal drives which.
    tr = []
    if spec.get("show_subtotal", True):
        tr.append(f"<tr><td>SUBTOTAL</td><td class=v>{cur}&nbsp;&nbsp;{money(spec['subtotal'])}</td></tr>")
        tr.append(f"<tr><td>{esc(freight_label)}</td><td class=v>"
                  f"{cur+'&nbsp;&nbsp;' if isinstance(freight_val,(int,float)) else ''}{freight_txt}</td></tr>")
    tr.append(f"<tr><td>{total_label}</td><td class=v>{cur}&nbsp;&nbsp;{money(spec['total'])}</td></tr>")
    totals_rows = "".join(tr)
    return f"""<!doctype html><html><head><meta charset=utf-8><style>
@page {{ size: Letter; margin: 14mm 12mm; }}
body {{ font-family: Calibri, Carlito, sans-serif; font-size: 10pt; color:#000; }}
.top {{ display:flex; align-items:flex-start; gap:10px; }}
.top img {{ height:42px; }}
h1 {{ font-size:19pt; margin:2px 0 0; text-align:center; flex:1; }}
.sub {{ font-size:15pt; font-weight:bold; text-align:right; margin:6px 0 10px; }}
.small {{ font-size:8.5pt; }}
.cols {{ display:flex; gap:0; }}
.left {{ width:52%; }}
.right {{ width:48%; }}
.box {{ border:1px solid #000; padding:4px 6px; }}
.box + .box {{ border-top:none; }}
.box b {{ display:block; }}
table.meta {{ width:100%; border-collapse:collapse; margin-top:2px; }}
table.meta td {{ padding:1px 4px; vertical-align:top; }}
td.ml {{ text-align:right; font-weight:bold; white-space:nowrap; }}
td.mv {{ border-left:1px solid #000; padding-left:5px; }}
table.items {{ width:100%; border-collapse:collapse; margin-top:8px; table-layout:fixed; }}
table.items th {{ border:1px solid #000; background:#fff; font-size:9.5pt; padding:2px; }}
table.items td {{ border-left:1px solid #000; border-right:1px solid #000; padding:1px 4px; font-size:10pt; }}
table.items tr:first-child td {{ padding-top:5px; }}
tr.blank td {{ height:15px; }}
table.items tr:last-child td {{ border-bottom:1px solid #000; }}
.c {{ text-align:center; }} .n {{ text-align:right; white-space:nowrap; }}
table.tot {{ border-collapse:collapse; margin-left:auto; margin-top:0; }}
table.tot td {{ padding:1px 5px; font-weight:bold; }}
table.tot td.v {{ border:1px solid #000; text-align:right; min-width:92px; }}
</style></head><body>
<div class=top>{f'<img src="{logo}">' if logo else ''}<h1>{AES_NAME}</h1></div>
<div class=small>{AES_ADDR}</div>
<div class=small>{AES_CONTACT}</div>
{f'<div class=small>{AES_FEI}</div>' if not oa else ''}
<div class=sub>{title}</div>
<div class=cols>
  <div class=left>
    <div class=box><b>{'SOLD TO:' if oa else 'TO:'}</b>{addr_block(spec['sold_to'])}</div>
    {f'<div class=box><b>SHIP TO:</b>{addr_block(ship)}</div>' if ship else ''}
  </div>
  <div class=right><table class=meta>{metarows}</table></div>
</div>
<table class=items>
<colgroup><col style="width:7%"><col style="width:8%">{'<col style="width:12%">' if not oa else ''}
<col><col style="width:14%"><col style="width:15%"></colgroup>
<tr><th>{'LINE' if not oa else 'ITEM'}</th><th>QTY</th>{hdr_item}<th>DESCRIPTION</th>
<th>UN PRICE</th><th>{hdr_total}</th></tr>
{''.join(rows)}
</table>
<table class=tot>
{totals_rows}
</table>
</body></html>"""


def render_pdf(spec, out_pdf):
    if not os.path.exists(CHROME):
        raise SystemExit("Google Chrome not found; cannot render PDF")
    with tempfile.TemporaryDirectory() as td:
        html = os.path.join(td, "doc.html")
        open(html, "w", encoding="utf-8").write(html_for(spec))
        subprocess.run([CHROME, "--headless", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={out_pdf}", html],
                       check=True, capture_output=True, timeout=120)
    return out_pdf


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec")
    ap.add_argument("--out", default=None, help="output dir (default: project folder in spec)")
    ap.add_argument("--no-xlsx", action="store_true")
    ap.add_argument("--no-pdf", action="store_true")
    a = ap.parse_args()
    spec = json.load(open(a.spec, encoding="utf-8"))

    errs, warns = validate(spec)
    for w in warns:
        print(f"  warning: {w}")
    if errs:
        print("VALIDATION FAILED — nothing written:")
        for e in errs:
            print("   ✗", e)
        sys.exit(1)
    print(f"  validated: {len(spec['items'])} lines, subtotal {money(spec['subtotal'])}, "
          f"total {money(spec['total'])}")

    outdir = a.out or spec["project_folder"]
    os.makedirs(outdir, exist_ok=True)

    if not a.no_xlsx and spec.get("workbook") and spec.get("cellmap"):
        wbk = spec["workbook"]
        backup = wbk + ".bak"
        if not os.path.exists(backup):
            shutil.copy2(wbk, backup)
            print(f"  backup: {os.path.basename(backup)}")
        write_tab(spec, wbk, spec["cellmap"]["tab"])
        print(f"  wrote tab '{spec['cellmap']['tab']}' in {os.path.basename(wbk)}")

    if not a.no_pdf:
        out = os.path.join(outdir, spec["pdf_name"])
        render_pdf(spec, out)
        print(f"  wrote PDF {spec['pdf_name']} ({os.path.getsize(out):,} bytes)")


if __name__ == "__main__":
    main()
